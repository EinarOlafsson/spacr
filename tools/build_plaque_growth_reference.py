"""Derive an experimental plaque-size anchor from Kelsen et al. 2023 S20 Data.

Run with the downloaded supplementary XLSX as the positional argument.
The paper and workbook are CC BY; retain attribution in generated outputs.
This evaluates endpoint repeatability, not a multi-time growth curve.
"""
import argparse
from collections import defaultdict
import csv
import hashlib
import json
import math
from pathlib import Path
import statistics


def build(source, destination):
    import openpyxl

    book = openpyxl.load_workbook(source, data_only=True)
    pooled = book['Pooled_Plaque#_Areapixels_mm^2']
    ratios = [pooled.cell(13, c + 10).value / pooled.cell(13, c).value for c in (4, 5, 6)]
    area_factor = statistics.median(ratios)
    if not all(math.isclose(r, area_factor, rel_tol=1e-9) for r in ratios):
        raise ValueError('The three physical-area conversion factors disagree.')
    groups = defaultdict(list)
    rows = []
    well = 0
    date = None
    for row in book['0uM KNX-002'].values:
        if isinstance(row[0], int) and row[1] == 'RH WT':
            date = str(row[0])
            well += 1
        if (isinstance(row[1], (int, float)) and isinstance(row[2], str)
                and row[2].endswith('.tif') and isinstance(row[3], (int, float))):
            if date is None or row[3] <= 0:
                raise ValueError('Unassigned or nonpositive plaque area.')
            diameter = math.sqrt(4 * row[3] * area_factor / math.pi) * 1000
            groups[date].append(diameter)
            rows.append(dict(experiment=date, well=well, area_px=row[3],
                             area_um2=row[3] * area_factor * 1e6, formation_hours=168))
    experiments = []
    for date, diameters in sorted(groups.items()):
        count = math.ceil(len(diameters) * .25)
        experiments.append(dict(experiment=date, count=len(diameters), selected=count,
                                diameter_um=statistics.median(sorted(diameters)[-count:])))
    anchor = statistics.median(e['diameter_um'] for e in experiments)
    for e in experiments:
        other = statistics.median(x['diameter_um'] for x in experiments if x is not e)
        e['held_out_estimated_hours'] = 168 * e['diameter_um'] / other
        e['held_out_absolute_error_hours'] = abs(e['held_out_estimated_hours'] - 168)
    reference = dict(
        id='kelsen2023_rh_hff_vehicle_day7_v1',
        attribution='Kelsen et al. (2023), PLOS Biology 21(5): e3002110, S20 Data. CC BY 4.0.',
        url='https://doi.org/10.1371/journal.pbio.3002110',
        data_url='https://journals.plos.org/plosbiology/article/file?id=10.1371%2Fjournal.pbio.3002110.s036&type=supplementary',
        sha256=hashlib.sha256(Path(source).read_bytes()).hexdigest(),
        selection='0uM KNX-002 sheet, WT columns B-D, RH WT experiment dates; mutants and treated wells excluded',
        physical_conversion='Median N13/D13, O13/E13, P13/F13 of Pooled_Plaque#_Areapixels_mm^2 sheet',
        area_mm2_per_pixel=area_factor, reference_hours=168,
        reference_diameter_um=anchor, selected_fraction=.25,
        experimental=True, temporal_validation=False,
        assumption='Equivalent diameter grows linearly from zero with elapsed time; no fitted lag. Growth can differ by strain, host, treatment and assay.',
        experiments=experiments,
        endpoint_mae_hours=statistics.mean(e['held_out_absolute_error_hours'] for e in experiments),
    )
    destination.mkdir(parents=True, exist_ok=True)
    (destination / 'reference.json').write_text(json.dumps(reference, indent=2) + '\n')
    with (destination / 'control_areas.csv').open('w', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    print(json.dumps(reference, indent=2))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path)
    parser.add_argument('--destination', type=Path, default=Path('tests/data/plaque_growth'))
    args = parser.parse_args()
    build(args.source, args.destination)
