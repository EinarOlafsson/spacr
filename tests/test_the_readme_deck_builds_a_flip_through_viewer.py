"""The README's info deck: one .pptx becomes pictures, a PDF and a viewer.

2026-09-21, the maintainer: "add it to the spacr readme and let people flip
through the slides as a kind of info deck". The rendering is LibreOffice's
and poppler's; what is tested here is what spaCR's own tool does with it.
"""
import importlib.util
import json
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
SPEC = importlib.util.spec_from_file_location(
    "build_readme_deck", ROOT / "tools" / "build_readme_deck.py")
deck = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(deck)


def _fake_render(pdf, folder, width, quality):
    from PIL import Image

    folder.mkdir(parents=True, exist_ok=True)
    out = []
    for number in (1, 2, 3):
        path = folder / f"slide_{number:02d}.jpg"
        Image.new("RGB", (width, width * 9 // 16), (number * 60, 20, 40)).save(path)
        out.append(path)
    return out


def test_a_deck_becomes_slides_a_pdf_and_a_viewer_that_needs_no_server(
        tmp_path, monkeypatch):
    pptx = tmp_path / "deck.pptx"
    pptx.write_bytes(b"not really a deck")
    monkeypatch.setattr(deck, "to_pdf", lambda p, work: work / "deck.pdf")
    monkeypatch.setattr(deck, "render", _fake_render)
    monkeypatch.setattr(deck, "animations", lambda p, folder: {})
    monkeypatch.setattr(deck, "titles",
                        lambda pdf, n: ["spaCR", "What is spaCR?", "Try it"])
    out = tmp_path / "out"
    assert deck.main([str(pptx), "--out", str(out), "--width", "160"]) == 0
    manifest = json.loads((out / "slides.json").read_text())
    assert manifest["count"] == 3
    assert [s["title"] for s in manifest["slides"]] == [
        "spaCR", "What is spaCR?", "Try it"]
    assert all((out / s["image"]).is_file() and (out / s["thumb"]).is_file()
               for s in manifest["slides"])
    assert (out / "spacr_deck.pdf").read_bytes().startswith(b"%PDF")
    page = (out / "index.html").read_text()
    assert "__SLIDES__" not in page, "the slide list is written into the page"
    assert '"What is spaCR?"' in page
    assert "ArrowRight" in page and "touchstart" in page


def test_the_viewer_loads_nothing_from_another_site():
    page = (ROOT / "tools" / "readme_deck_viewer.html").read_text()
    assert "<script src" not in page and "<link" not in page
    assert page.count("__SLIDES__") == 1


def test_an_animated_gif_is_read_out_of_the_deck_with_its_place(tmp_path):
    """A picture export keeps only a GIF's first frame, so the viewer plays
    the GIF itself over the slide, where the deck put it."""
    pptx = pytest.importorskip("pptx")
    from PIL import Image
    from pptx.util import Inches

    frames = [Image.new("RGB", (40, 40), (i * 20, 0, 0)) for i in range(4)]
    gif = tmp_path / "moving.gif"
    frames[0].save(gif, save_all=True, append_images=frames[1:], duration=50)
    deck_file = tmp_path / "deck.pptx"
    presentation = pptx.Presentation()
    presentation.slide_width, presentation.slide_height = Inches(10), Inches(5)
    presentation.slides.add_slide(presentation.slide_layouts[6])
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(str(gif), Inches(5), Inches(1), Inches(2), Inches(2))
    presentation.save(deck_file)

    found = deck.animations(deck_file, tmp_path / "anim")
    assert list(found) == [2], "only the second slide has an animation"
    (moving,) = found[2]
    assert (moving["x"], moving["y"], moving["w"], moving["h"]) == \
        pytest.approx((0.5, 0.2, 0.2, 0.4), abs=1e-3)
    copied = tmp_path / moving["src"]
    assert Image.open(copied).n_frames == 4


def test_every_github_page_links_back_and_next_and_wraps(tmp_path):
    pages = deck.github_pages(tmp_path, ["One", "Two", "Three"])
    text = [p.read_text() for p in pages]
    assert 'href="03.md">← Back' in text[0] and 'href="02.md">Next →' in text[0]
    assert 'href="02.md">← Back' in text[2] and 'href="01.md">Next →' in text[2]
    assert 'src="../slides/slide_02.jpg"' in text[1]


def test_the_title_slide_is_restamped_with_the_version(tmp_path):
    """2026-09-21, the maintainer: "the title page should autoupdate with
    version bumps"."""
    from PIL import Image

    for folder in ("slides", "thumbs"):
        (tmp_path / folder).mkdir()
    Image.new("RGB", (800, 450), "black").save(tmp_path / "title_base.jpg")
    Image.new("RGB", (800, 450), "black").save(tmp_path / "slides" / "slide_01.jpg")
    (tmp_path / "slides.json").write_text(json.dumps({
        "slides": [{"image": "slides/slide_01.jpg", "thumb": "thumbs/slide_01.jpg"}],
        "title_line": {"x": 0.05, "y": 0.85, "w": 0.7, "h": 0.06, "size": 0.04,
                       "colour": "#FFFFFF", "text": "spaCR {version}  ·  x"}}))
    assert deck.stamp_title(tmp_path, "9.8.7.6") == "spaCR 9.8.7.6  ·  x"
    drawn = Image.open(tmp_path / "slides" / "slide_01.jpg").convert("L")
    assert drawn.crop((0, 360, 800, 420)).getextrema()[1] > 128
    assert json.loads((tmp_path / "slides.json").read_text())["version"] == "9.8.7.6"
    assert (tmp_path / "spacr_deck.pdf").is_file()


def test_the_published_deck_carries_the_current_version():
    manifest = json.loads((ROOT / "docs" / "source" / "_static" / "deck"
                           / "slides.json").read_text(encoding="utf-8"))
    assert manifest["version"] == deck.current_version(), (
        "run python tools/build_readme_deck.py --stamp "
        "(packaging/release.py does it on every bump)")
